# -*- coding: utf-8 -*-
"""Word building blocks: headings, body text, shaded example boxes, tables."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor as C
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK=C(0x1c,0x1f,0x23); INK2=C(0x46,0x4b,0x52); MUT=C(0x6f,0x75,0x7e)
TEAL=C(0x1f,0x6f,0x6a); PLUM=C(0x8a,0x3d,0x5e); AMBER=C(0x8a,0x64,0x10); RED=C(0x9b,0x3b,0x2f)
SANS="Calibri"; MONO="Consolas"

def shade(cell, hexfill):
    el = OxmlElement("w:shd"); el.set(qn("w:val"),"clear"); el.set(qn("w:fill"),hexfill)
    cell._tc.get_or_add_tcPr().append(el)

def no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top","left","bottom","right","insideH","insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"),"none"); x.set(qn("w:sz"),"0")
        b.append(x)
    tblPr.append(b)

def doc_new():
    d = Document()
    for name,size,color,bold,space_before,space_after in [
        ("Normal",10.5,INK2,False,0,6),
    ]:
        st = d.styles[name]; st.font.name=SANS; st.font.size=Pt(size)
        st.font.color.rgb=color
        st.paragraph_format.space_after=Pt(space_after)
        st.paragraph_format.line_spacing=1.15
    for s in d.sections:
        s.top_margin=Inches(0.85); s.bottom_margin=Inches(0.85)
        s.left_margin=Inches(0.9); s.right_margin=Inches(0.9)
    return d

def h(d, text, level=1):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(20 if level==1 else 14)
    pf.space_after  = Pt(6 if level==1 else 4)
    pf.keep_with_next = True
    r = p.add_run(text); r.bold=True; r.font.name=SANS
    r.font.size = Pt(16 if level==1 else 12.5)
    r.font.color.rgb = INK if level==1 else TEAL
    return p

def para(d, text, size=10.5, color=INK2, bold=False, italic=False,
         space_after=6, mono=False, indent=0):
    p = d.add_paragraph(); p.paragraph_format.space_after=Pt(space_after)
    if indent: p.paragraph_format.left_indent=Inches(indent)
    for chunk, is_bold in _split_bold(text):
        r = p.add_run(chunk); r.bold = bold or is_bold; r.italic = italic
        r.font.name = MONO if mono else SANS
        r.font.size = Pt(size); r.font.color.rgb = color
    return p

def _split_bold(text):
    """**bold** inside a string."""
    out=[]; i=0
    while True:
        a = text.find("**", i)
        if a == -1: out.append((text[i:], False)); break
        b = text.find("**", a+2)
        if b == -1: out.append((text[i:], False)); break
        if a > i: out.append((text[i:a], False))
        out.append((text[a+2:b], True)); i = b+2
    return [(t,bd) for t,bd in out if t]

def bullet(d, text, size=10.5):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Inches(0.28)
    for chunk, bd in _split_bold(text):
        r=p.add_run(chunk); r.bold=bd; r.font.name=SANS
        r.font.size=Pt(size); r.font.color.rgb=INK2
    return p

def box(d, label, lines, fill="F2F6F5", accent=TEAL):
    """A shaded single-cell panel for a worked example."""
    t = d.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT
    no_borders(t); c = t.rows[0].cells[0]; shade(c, fill)
    c.width = Inches(6.7)
    first = c.paragraphs[0]
    first.paragraph_format.space_after = Pt(3)
    r = first.add_run(label); r.bold=True; r.font.size=Pt(9.5)
    r.font.color.rgb = accent; r.font.name=SANS
    for ln in lines:
        mono = ln.startswith("`")
        txt = ln.strip("`")
        p = c.add_paragraph(); p.paragraph_format.space_after=Pt(2)
        for chunk, bd in _split_bold(txt):
            rr = p.add_run(chunk); rr.bold=bd
            rr.font.name = MONO if mono else SANS
            rr.font.size = Pt(9.5 if mono else 10)
            rr.font.color.rgb = INK if mono else INK2
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def table(d, headers, rows, widths=None, size=9.5, header_fill="1C1F23"):
    t = d.add_table(rows=1, cols=len(headers)); t.style="Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i,htxt in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, header_fill)
        p = c.paragraphs[0]; p.paragraph_format.space_after=Pt(1)
        r = p.add_run(htxt); r.bold=True; r.font.size=Pt(8.5)
        r.font.color.rgb=C(0xff,0xff,0xff); r.font.name=SANS
    for ri,row in enumerate(rows):
        cells = t.add_row().cells
        for i,val in enumerate(row):
            if ri % 2 == 1: shade(cells[i], "F7F7F6")
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(1)
            mono = isinstance(val,str) and val.startswith("`")
            txt = val.strip("`") if isinstance(val,str) else str(val)
            for chunk, bd in _split_bold(txt):
                r = p.add_run(chunk); r.bold=bd
                r.font.name = MONO if mono else SANS
                r.font.size=Pt(size); r.font.color.rgb=INK2
    if widths:
        for i,w in enumerate(widths):
            for row in t.rows: row.cells[i].width = Inches(w)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def fig(d, path, width=6.6, caption=None):
    d.add_picture(path, width=Inches(width))
    d.paragraphs[-1].alignment = AL.CENTER
    if caption:
        p = d.add_paragraph(); p.alignment=AL.CENTER
        p.paragraph_format.space_after=Pt(10)
        r=p.add_run(caption); r.italic=True; r.font.size=Pt(8.5)
        r.font.color.rgb=MUT; r.font.name=SANS
